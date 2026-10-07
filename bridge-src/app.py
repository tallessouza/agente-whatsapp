import asyncio
import base64
import hashlib
from io import BytesIO
import json
import logging
import os
import re
import sqlite3
import time
import uuid
import handover
import sdr
import tempfile
import threading
from zipfile import ZipFile
from xml.etree import ElementTree
from collections import defaultdict
from json import dumps as json_dumps
from mimetypes import guess_extension, guess_type
from pathlib import Path
from typing import Any

import httpx
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse

DB_PATH = Path(os.getenv("BRIDGE_DB", "/data/bridge.db"))
EVOLUTION_URL = os.getenv("EVOLUTION_URL", "http://evolution:4000").rstrip("/")
EVOLUTION_API_KEY = os.getenv("EVOLUTION_API_KEY", "")
EVOLUTION_INSTANCE = os.getenv("EVOLUTION_INSTANCE", "")
EVOLUTION_SEND_PATH = os.getenv("EVOLUTION_SEND_PATH", "/send/text")
EVOLUTION_SEND_MEDIA_PATH = os.getenv("EVOLUTION_SEND_MEDIA_PATH", "/send/media")
BRIDGE_WEBHOOK_TOKEN = os.getenv("BRIDGE_WEBHOOK_TOKEN", "")
CHATWOOT_URL = os.getenv("CHATWOOT_URL", "http://chatwoot_rails:3000").rstrip("/")
CHATWOOT_INBOX_IDENTIFIER = os.getenv("CHATWOOT_INBOX_IDENTIFIER", "")
CHATWOOT_TOKEN = os.getenv("CHATWOOT_API_ACCESS_TOKEN", "")
CHATWOOT_ACCOUNT_ID = os.getenv("CHATWOOT_ACCOUNT_ID", "")
CHATWOOT_APP_TOKEN = os.getenv("CHATWOOT_APP_TOKEN", "")
CHATWOOT_INBOX_ID = os.getenv("CHATWOOT_INBOX_ID", "1")
HERMES_URL = os.getenv("HERMES_URL", "http://hermes:8642/v1").rstrip("/")
HERMES_KEY = os.getenv("HERMES_API_KEY", "")
HERMES_MODEL = os.getenv("HERMES_MODEL", "hermes-agent")
# Hermes freezes the system prompt per session: bump the prefix to roll out a new persona.
# v13 = SOUL with SPICED and the two offers (06/10/2026). The service env wins over this default.
HERMES_SESSION_PREFIX = os.getenv("HERMES_SESSION_PREFIX", "wa:v13:")
# Follow-up sending window (local hours, end exclusive); 0-24 sends around the clock.
SDR_BUSINESS_START_HOUR = int(os.getenv("SDR_BUSINESS_START_HOUR", "9"))
SDR_BUSINESS_END_HOUR = int(os.getenv("SDR_BUSINESS_END_HOUR", "20"))
HERMES_MEDIA_MAX_BYTES = int(os.getenv("HERMES_MEDIA_MAX_BYTES", str(5 * 1024 * 1024)))
HERMES_MEDIA_ROOT = Path(os.getenv("HERMES_MEDIA_ROOT", "/hermes-data")).resolve()
HERMES_MEDIA_PREFIX = os.getenv("HERMES_MEDIA_PREFIX", "/opt/data").rstrip("/")
SDR_BRAIN_URL = os.getenv("SDR_BRAIN_URL", "").rstrip("/")
SDR_BRAIN_TOKEN = os.getenv("SDR_BRAIN_TOKEN", "")
SDR_REVIEW_TIMEOUT = float(os.getenv("SDR_REVIEW_TIMEOUT", "12"))
HERMES_MESSAGE_DEBOUNCE_SECONDS = float(os.getenv("HERMES_MESSAGE_DEBOUNCE_SECONDS", "10"))
HERMES_MEDIA_HISTORY_SECONDS = int(os.getenv("HERMES_MEDIA_HISTORY_SECONDS", "600"))
EVOLUTION_MAX_MESSAGE_LENGTH = int(os.getenv("EVOLUTION_MAX_MESSAGE_LENGTH", "4096"))
BRIDGE_GROUP_POLICY = os.getenv("BRIDGE_GROUP_POLICY", "disabled").strip().lower()
BRIDGE_GROUP_ALLOW_FROM = {item.strip() for item in os.getenv("BRIDGE_GROUP_ALLOW_FROM", "").split(",") if item.strip()}
BRIDGE_REQUIRE_MENTION = os.getenv("BRIDGE_REQUIRE_MENTION", "false").lower() in {"1", "true", "yes", "on"}
BRIDGE_MENTION_PATTERNS = [item.strip() for item in os.getenv("BRIDGE_MENTION_PATTERNS", "").split(",") if item.strip()]
STT_ENABLED = os.getenv("BRIDGE_STT_ENABLED", "true").lower() not in {"0", "false", "no"}
STT_MODEL = os.getenv("BRIDGE_STT_MODEL", "small")
STT_LANGUAGE = os.getenv("BRIDGE_STT_LANGUAGE", "pt") or None
OCR_ENABLED = os.getenv("BRIDGE_OCR_ENABLED", "true").lower() not in {"0", "false", "no"}
OCR_LANGUAGE = os.getenv("BRIDGE_OCR_LANGUAGE", "por+eng") or "por+eng"
MEDIA_DIR = DB_PATH.parent / "media"

app = FastAPI(title="Evolution Go ↔ Chatwoot bridge")
logger = logging.getLogger("bridge")
conversation_lock = asyncio.Lock()
HERMES_MAX_CONCURRENCY = max(1, int(os.getenv("HERMES_MAX_CONCURRENCY", "8")))
HERMES_429_RETRY_SECONDS = float(os.getenv("HERMES_429_RETRY_SECONDS", "180"))
hermes_slots = asyncio.Semaphore(HERMES_MAX_CONCURRENCY)
hermes_pending: dict[str, list[tuple[str, list[tuple[bytes, str, str]] | None, str, int, str, dict[str, Any] | None]]] = defaultdict(list)
hermes_tasks: dict[str, asyncio.Task] = {}
_stt_model: Any = None
_stt_lock = threading.Lock()
_outbox_task: asyncio.Task | None = None
_followup_task: asyncio.Task | None = None


@app.on_event("startup")
async def start_background_workers() -> None:
    global _outbox_task, _followup_task
    db().close()
    _outbox_task = asyncio.create_task(outbox_worker(), name="evolution-outbox")
    _followup_task = asyncio.create_task(followup_worker(), name="sdr-followups")


@app.on_event("shutdown")
async def stop_background_workers() -> None:
    global _outbox_task, _followup_task
    if _outbox_task:
        _outbox_task.cancel()
        try:
            await _outbox_task
        except asyncio.CancelledError:
            pass
        _outbox_task = None
    if _followup_task:
        _followup_task.cancel()
        try: await _followup_task
        except asyncio.CancelledError: pass
        _followup_task = None


def db() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("CREATE TABLE IF NOT EXISTS contacts (jid TEXT PRIMARY KEY, source_id TEXT NOT NULL, conversation_id INTEGER NOT NULL, internal_id INTEGER)")
    try:
        conn.execute("ALTER TABLE contacts ADD COLUMN internal_id INTEGER")
    except sqlite3.OperationalError:
        pass
    conn.execute("CREATE TABLE IF NOT EXISTS messages (external_id TEXT PRIMARY KEY, chatwoot_id INTEGER, created_at REAL NOT NULL DEFAULT 0)")
    try:
        conn.execute("ALTER TABLE messages ADD COLUMN created_at REAL NOT NULL DEFAULT 0")
    except sqlite3.OperationalError:
        pass
    conn.execute("UPDATE messages SET created_at=? WHERE created_at=0", (time.time(),))
    conn.execute("DELETE FROM messages WHERE created_at < ?", (time.time() - 86400,))
    conn.execute("CREATE TABLE IF NOT EXISTS event_history (external_id TEXT PRIMARY KEY, event_type TEXT NOT NULL, jid TEXT, target_id TEXT, payload TEXT NOT NULL, created_at REAL NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS chatwoot_deliveries (external_id TEXT PRIMARY KEY)")
    conn.execute("CREATE TABLE IF NOT EXISTS media_history (external_id TEXT PRIMARY KEY, jid TEXT NOT NULL, path TEXT NOT NULL, mimetype TEXT NOT NULL, filename TEXT NOT NULL, created_at REAL NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS outbox (id INTEGER PRIMARY KEY AUTOINCREMENT, kind TEXT NOT NULL, number TEXT NOT NULL, payload TEXT NOT NULL, attempts INTEGER NOT NULL DEFAULT 0, next_attempt REAL NOT NULL, status TEXT NOT NULL DEFAULT 'pending', last_error TEXT, created_at REAL NOT NULL)")
    conn.execute("UPDATE outbox SET status='pending' WHERE status='inflight'")
    conn.commit()
    return conn


def unwrap_message(message: dict[str, Any]) -> dict[str, Any]:
    """Unwrap common Baileys/Evolution envelopes before inspecting a message."""
    current = message if isinstance(message, dict) else {}
    for _ in range(5):
        wrapper = next((current.get(key) for key in (
            "ephemeralMessage", "viewOnceMessage", "viewOnceMessageV2", "viewOnceMessageV2Extension",
            "documentWithCaptionMessage", "deviceSentMessage",
        ) if isinstance(current.get(key), dict)), None)
        if not wrapper:
            break
        current = wrapper.get("message") or wrapper.get("Message") or wrapper
    return current


def text_from_message(message: dict[str, Any]) -> str:
    # WhatsApp template/interactive envelopes carry their visible text one or
    # two levels below the top-level message. Flatten the user-visible choice
    # so Chatwoot and Hermes see the same event as a typed reply.
    message = unwrap_message(message)
    payment = payment_event(message)
    if payment:
        amount = payment.get("amount")
        currency = payment.get("currency") or ""
        amount_text = f"{currency} {amount:.2f}" if isinstance(amount, (int, float)) else str(amount or "valor não informado")
        return f"[Solicitação de pagamento: {amount_text}]"
    location = message.get("locationMessage") or message.get("liveLocationMessage")
    if isinstance(location, dict):
        lat = location.get("degreesLatitude") or location.get("latitude")
        lon = location.get("degreesLongitude") or location.get("longitude")
        if lat is not None and lon is not None:
            label = location.get("name") or location.get("address") or "localização"
            return f"[Localização: {label} ({lat}, {lon}) https://maps.google.com/?q={lat},{lon}]"
    poll = message.get("pollCreationMessage") or message.get("pollCreationMessageV3")
    if isinstance(poll, dict):
        question = poll.get("name") or poll.get("question") or "enquete"
        options = [item.get("optionName") for item in poll.get("options", []) if isinstance(item, dict) and item.get("optionName")]
        return f"[Enquete: {question}" + (f" | Opções: {', '.join(options)}" if options else "") + "]"
    poll_update = message.get("pollUpdateMessage")
    if isinstance(poll_update, dict):
        return "[Voto de enquete recebido]"
    reaction = message.get("reactionMessage")
    if isinstance(reaction, dict):
        emoji = reaction.get("text") or reaction.get("emoji") or ""
        target = reaction.get("key", {}).get("id") if isinstance(reaction.get("key"), dict) else ""
        return f"[Reação recebida: {emoji}{f' na mensagem {target}' if target else ''}]"
    protocol = message.get("protocolMessage")
    if isinstance(protocol, dict):
        edited = protocol.get("editedMessage") or protocol.get("editedMessageV2")
        kind = str(protocol.get("type") or protocol.get("editedMessage") and "edit").lower()
        target = protocol.get("key", {}).get("id") if isinstance(protocol.get("key"), dict) else ""
        if isinstance(edited, dict):
            edited_text = text_from_message(edited)
            return f"[Mensagem editada{f' {target}' if target else ''}: {edited_text}]"
        label = "Mensagem revogada" if "revoke" in kind or kind in {"0", "5"} else "Mensagem atualizada"
        return f"[{label}{f': {target}' if target else ''}]"
    if isinstance(message.get("templateButtonReplyMessage"), dict):
        reply = message["templateButtonReplyMessage"]
        return str(reply.get("selectedDisplayText") or reply.get("selectedId") or "[botão selecionado]")
    if isinstance(message.get("buttonsResponseMessage"), dict):
        reply = message["buttonsResponseMessage"]
        return str(reply.get("selectedDisplayText") or reply.get("selectedButtonId") or "[botão selecionado]")
    if isinstance(message.get("listResponseMessage"), dict):
        reply = message["listResponseMessage"]
        return str(reply.get("title") or reply.get("singleSelectReply", {}).get("selectedRowId") or "[item selecionado]")
    if isinstance(message.get("interactiveResponseMessage"), dict):
        reply = message["interactiveResponseMessage"]
        native = reply.get("nativeFlowResponseMessage") or {}
        return str(native.get("paramsJson") or reply.get("body", {}).get("text") or "[interação recebida]")
    for key in ("conversation", "Conversation", "extendedTextMessage", "ExtendedTextMessage", "text", "Text", "caption", "Caption"):
        value = message.get(key)
        if isinstance(value, str):
            return value
        if isinstance(value, dict) and isinstance(value.get("text"), str):
            return value["text"]
    template = message.get("templateMessage")
    if isinstance(template, dict):
        for envelope in (template.get("hydratedTemplate"), template.get("fourRowTemplate"), template.get("hydratedFourRowTemplate")):
            if isinstance(envelope, dict):
                body = envelope.get("hydratedContentText") or envelope.get("contentText") or envelope.get("contentText")
                if body:
                    return str(body)
        return "[template recebido]"
    for key, value in message.items():
        if isinstance(value, dict) and isinstance(value.get("caption"), str):
            return value["caption"]
    return "[mídia recebida]"


def payment_event(message: dict[str, Any], *, key: dict[str, Any] | None = None, data: dict[str, Any] | None = None) -> dict[str, Any] | None:
    """Normalize WhatsApp's request-payment envelope without executing it.

    Evolution forwards Baileys' ``requestPaymentMessage`` as a normal Message
    event.  The bridge keeps the provider fields intact and adds transport
    context so Hermes can reason about the request as data rather than treating
    the green card as an opaque attachment.
    """
    message = unwrap_message(message)
    if not isinstance(message, dict):
        return None
    envelope_name = next((name for name in ("requestPaymentMessage", "paymentRequestMessage") if isinstance(message.get(name), dict)), None)
    raw: dict[str, Any] | None = message.get(envelope_name) if envelope_name else None
    if not raw and isinstance(message.get("interactiveMessage"), dict):
        interactive = message["interactiveMessage"]
        candidates: list[dict[str, Any]] = []
        native = interactive.get("nativeFlowMessage") or {}
        blobs = [native.get("messageParamsJson")]
        blobs.extend(button.get("buttonParamsJson") for button in native.get("buttons", []) if isinstance(button, dict))
        for blob in blobs:
            parsed: Any = blob
            if isinstance(blob, str):
                try:
                    parsed = json.loads(blob)
                except json.JSONDecodeError:
                    continue
            def collect(value: Any) -> None:
                if isinstance(value, dict):
                    marker = str(value.get("type", "")).lower()
                    keys = {str(key).lower() for key in value}
                    if marker in {"pix", "payment_info", "payment"} or any("pix" in key or "payment" in key for key in keys):
                        nested = next((item for name, item in value.items() if str(name).lower() in {"pix", "payment_info", "payment"} and isinstance(item, dict)), None)
                        candidates.append(nested or value)
                    for item in value.values():
                        collect(item)
                elif isinstance(value, list):
                    for item in value:
                        collect(item)
            collect(parsed)
        raw = candidates[0] if candidates else None
        if raw:
            envelope_name = "interactiveMessage.payment_info"
    if not raw:
        return None
    amount1000 = raw.get("amount1000") or raw.get("priceAmount1000")
    amount = None
    if isinstance(amount1000, (int, float)) and not isinstance(amount1000, bool):
        amount = amount1000 / 1000
    elif isinstance(amount1000, str):
        try:
            amount = int(amount1000) / 1000
        except ValueError:
            pass
    money = raw.get("amount") if isinstance(raw.get("amount"), dict) else {}
    if amount is None:
        candidate = money.get("value") or money.get("amount") or money.get("value1000")
        if isinstance(candidate, (int, float)) and not isinstance(candidate, bool):
            amount = candidate / 1000 if money.get("value1000") is not None else candidate
    currency = raw.get("currencyCodeIso4217") or raw.get("currency") or money.get("currencyCodeIso4217") or money.get("currency")
    context = {
        "message_id": (key or {}).get("id"),
        "remote_jid": (key or {}).get("remoteJid"),
        "from_me": bool((key or {}).get("fromMe")),
        "push_name": (data or {}).get("pushName") or (data or {}).get("PushName"),
    }
    context = {name: value for name, value in context.items() if value not in (None, "")}
    return {
        "event_type": "payment_request",
        "provider": "whatsapp",
        "provider_message_type": envelope_name,
        "payment_type": raw.get("type") or ("pix" if envelope_name == "interactiveMessage.payment_info" else "request_payment"),
        "amount": amount,
        "amount1000": amount1000,
        "currency": currency,
        "request_from": raw.get("requestFrom"),
        "pix_name": raw.get("name"),
        "pix_key_type": raw.get("keyType"),
        "pix_key": raw.get("key") or raw.get("copyCode"),
        "expiry_timestamp": raw.get("expiryTimestamp"),
        "note": text_from_message(raw.get("noteMessage") or {}) if isinstance(raw.get("noteMessage"), dict) else None,
        "provider_payload": raw,
        "context": context,
        "execution": {"authorized": False, "executed": False},
    }


def has_media_envelope(message: dict[str, Any]) -> bool:
    message = unwrap_message(message)
    return isinstance(message, dict) and any(isinstance(message.get(key), dict) for key in ("imageMessage", "videoMessage", "audioMessage", "documentMessage", "stickerMessage"))


def is_state_event(message: dict[str, Any]) -> bool:
    """Events that belong in history but must not start an agent turn."""
    message = unwrap_message(message)
    return any(isinstance(message.get(key), dict) for key in ("reactionMessage", "protocolMessage", "pollUpdateMessage"))


def media_info(message: dict[str, Any]) -> tuple[bytes, str, str] | None:
    """Extract Evolution Go's processed media payload for Chatwoot uploads."""
    message = unwrap_message(message)
    if not isinstance(message, dict):
        return None
    media_key = next((k for k in ("imageMessage", "videoMessage", "audioMessage", "documentMessage", "stickerMessage") if isinstance(message.get(k), dict)), None)
    if not media_key:
        return None
    media = message[media_key]
    raw = media.get("base64") or message.get("base64")
    mimetype = str(media.get("mimetype") or message.get("mimetype") or "application/octet-stream")
    if isinstance(raw, str) and raw:
        encoded = raw.split(",", 1)[1] if raw.startswith("data:") and "," in raw else raw
        try:
            payload = base64.b64decode(encoded)
        except (ValueError, base64.binascii.Error):
            payload = b""
        if payload:
            extension = guess_extension(mimetype) or ".bin"
            prefix = {"imageMessage": "image", "stickerMessage": "sticker", "videoMessage": "video", "audioMessage": "audio", "documentMessage": "document"}.get(media_key, "media")
            filename = str(media.get("fileName") or media.get("filename") or f"{prefix}{extension}")
            if "." not in filename.rsplit("/", 1)[-1]:
                filename += extension
            return payload, mimetype, filename
    return None


def quoted_context(message: dict[str, Any]) -> tuple[str, str, str]:
    """Extract Evolution's quoted-message identity, visible text and participant."""
    message = unwrap_message(message)
    candidates = [message.get("messageContextInfo"), message.get("contextInfo")]
    for value in message.values():
        if isinstance(value, dict):
            candidates.append(value.get("contextInfo"))
    context = next((value for value in candidates if isinstance(value, dict)), {})
    quoted = context.get("quotedMessage") or context.get("quotedMessageV2") or {}
    quoted = unwrap_message(quoted) if isinstance(quoted, dict) else {}
    quote_text = text_from_message(quoted) if quoted else ""
    if quote_text == "[mídia recebida]":
        quote_text = ""
    return (
        str(context.get("stanzaId") or context.get("quotedMessageId") or ""),
        quote_text,
        str(context.get("participant") or context.get("remoteJid") or ""),
    )


def quoted_media_from_cache(message: dict[str, Any], quoted_id: str) -> list[tuple[bytes, str, str]]:
    """Resolve media from an inline quoted envelope or the persistent media cache."""
    message = unwrap_message(message)
    candidates = [message.get("messageContextInfo"), message.get("contextInfo")]
    for value in message.values():
        if isinstance(value, dict):
            candidates.append(value.get("contextInfo"))
    context = next((value for value in candidates if isinstance(value, dict)), {})
    quoted = context.get("quotedMessage") or context.get("quotedMessageV2") or {}
    if isinstance(quoted, dict):
        inline = media_info(quoted)
        if inline:
            return [inline]
    if not quoted_id:
        return []
    conn = db()
    rows = conn.execute("SELECT path,mimetype,filename FROM media_history WHERE external_id=?", (quoted_id,)).fetchall()
    conn.close()
    result = []
    for path, mimetype, filename in rows:
        try:
            result.append((Path(path).read_bytes(), mimetype, filename))
        except OSError:
            logger.warning("Quoted media cache missing id=%s path=%s", quoted_id, path)
    return result


def _transcribe_audio_sync(payload: bytes, mimetype: str, filename: str) -> str:
    """Transcribe one WhatsApp voice note locally with faster-whisper."""
    global _stt_model
    if not STT_ENABLED or not mimetype.startswith("audio/"):
        return ""
    from faster_whisper import WhisperModel
    with _stt_lock:
        if _stt_model is None:
            logger.info("Loading local STT model=%s", STT_MODEL)
            _stt_model = WhisperModel(STT_MODEL, device="cpu", compute_type="int8")
    suffix = Path(filename).suffix
    if not suffix or suffix == ".bin":
        # Evolution often names voice notes audio.bin while retaining the MIME.
        suffix = {
            "audio/ogg": ".ogg",
            "audio/ogg; codecs=opus": ".ogg",
            "audio/webm": ".webm",
            "audio/mp4": ".m4a",
            "audio/mpeg": ".mp3",
            "audio/wav": ".wav",
        }.get(mimetype.lower(), ".ogg")
    with tempfile.NamedTemporaryFile(prefix="bridge-stt-", suffix=suffix, delete=False) as tmp:
        tmp.write(payload)
        path = tmp.name
    try:
        segments, _info = _stt_model.transcribe(path, language=STT_LANGUAGE, vad_filter=True)
        return " ".join(segment.text.strip() for segment in segments if segment.text.strip()).strip()
    finally:
        try:
            Path(path).unlink()
        except OSError:
            pass


async def transcribe_audio_media(media: list[tuple[bytes, str, str]]) -> list[str]:
    audio = [item for item in media if item[1].startswith("audio/")]
    if not audio:
        return []
    results = []
    for payload, mimetype, filename in audio:
        try:
            transcript = await asyncio.to_thread(_transcribe_audio_sync, payload, mimetype, filename)
            if transcript:
                results.append(transcript)
        except Exception:
            logger.exception("Local audio transcription failed filename=%s", filename)
    return results


def extract_document_media(media: list[tuple[bytes, str, str]]) -> list[str]:
    """Extract text from supported document attachments for the Hermes turn."""
    results = []
    for payload, mimetype, filename in media:
        normalized = mimetype.split(";", 1)[0].lower()
        try:
            text = ""
            if normalized == "application/pdf":
                from pypdf import PdfReader
                reader = PdfReader(BytesIO(payload))
                text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            elif normalized in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"}:
                from docx import Document
                document = Document(BytesIO(payload))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()).strip()
            elif normalized in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"}:
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
                rows = []
                for sheet in workbook.worksheets:
                    rows.append(f"[Planilha: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value) for value in row if value is not None]
                        if values:
                            rows.append(" | ".join(values))
                text = "\n".join(rows).strip()
            elif normalized.startswith("text/") or normalized in {"application/json", "application/xml", "application/rtf"}:
                text = payload.decode("utf-8", errors="replace").strip()
            elif normalized == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
                from pptx import Presentation
                presentation = Presentation(BytesIO(payload))
                text = "\n".join(
                    shape.text for slide in presentation.slides for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ).strip()
            elif normalized in {"application/vnd.oasis.opendocument.text", "application/vnd.oasis.opendocument.spreadsheet"}:
                with ZipFile(BytesIO(payload)) as archive:
                    xml = archive.read("content.xml")
                root = ElementTree.fromstring(xml)
                text = "\n".join(item.text.strip() for item in root.iter() if item.text and item.text.strip()).strip()
            else:
                continue
            if text:
                # Keep a single attachment from consuming the whole context.
                results.append(f"[Texto extraído do documento {filename}]:\n{text[:24000]}")
            elif normalized == "application/pdf":
                ocr_text = ocr_pdf_sync(payload) if OCR_ENABLED else ""
                if ocr_text:
                    results.append(f"[OCR do documento {filename}]:\n{ocr_text[:24000]}")
                else:
                    logger.warning("PDF has no extractable text filename=%s", filename)
        except Exception:
            logger.exception("Document text extraction failed filename=%s", filename)
    return results


def ocr_pdf_sync(payload: bytes) -> str:
    """OCR scanned PDF pages locally; bounded to avoid blocking the bridge indefinitely."""
    if not OCR_ENABLED:
        return ""
    try:
        try:
            import fitz
        except ImportError:
            import pymupdf as fitz
        import pytesseract
        from PIL import Image
        document = fitz.open(stream=payload, filetype="pdf")
        pages = []
        for index, page in enumerate(document):
            if index >= 10:
                break
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            text = pytesseract.image_to_string(image, lang=OCR_LANGUAGE).strip()
            if text:
                pages.append(text)
        return "\n\n".join(pages).strip()
    except Exception:
        logger.exception("Scanned PDF OCR failed")
        return ""


def remember_media(jid: str, external_id: str, media: tuple[bytes, str, str]) -> None:
    """Persist a bounded-reference cache so a later text can refer to old media."""
    payload, mimetype, filename = media
    MEDIA_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix or (guess_extension(mimetype) or ".bin")
    digest = hashlib.sha256(f"{jid}:{external_id}".encode()).hexdigest()
    path = MEDIA_DIR / f"{digest}{suffix}"
    path.write_bytes(payload)
    conn = db()
    conn.execute("INSERT OR REPLACE INTO media_history VALUES (?,?,?,?,?,?)", (external_id, jid, str(path), mimetype, filename, time.time()))
    conn.commit()
    conn.close()


def recent_local_media(jid: str, limit: int = 3) -> list[tuple[bytes, str, str]]:
    cutoff = time.time() - HERMES_MEDIA_HISTORY_SECONDS
    conn = db()
    rows = conn.execute("SELECT path,mimetype,filename FROM media_history WHERE jid=? AND created_at>=? ORDER BY created_at DESC LIMIT ?", (jid, cutoff, limit)).fetchall()
    conn.close()
    result = []
    for path, mimetype, filename in rows:
        try:
            result.append((Path(path).read_bytes(), mimetype, filename))
        except OSError:
            continue
    return result


def jid_number(jid: str) -> str:
    return re.sub(r"[^0-9]", "", jid.split("@", 1)[0])


def group_is_allowed(jid: str) -> bool:
    """Apply the WhatsApp adapter's disabled/open/allowlist group policy."""
    if not jid.endswith("@g.us"):
        return True
    if BRIDGE_GROUP_POLICY == "open":
        return True
    if BRIDGE_GROUP_POLICY in {"allowlist", "allowed"}:
        return jid in BRIDGE_GROUP_ALLOW_FROM or jid_number(jid) in BRIDGE_GROUP_ALLOW_FROM
    return False


def group_triggered(message: dict[str, Any], content: str) -> bool:
    if not BRIDGE_REQUIRE_MENTION:
        return True
    message = unwrap_message(message)
    contexts = [message.get("messageContextInfo"), message.get("contextInfo")]
    for value in message.values():
        if isinstance(value, dict):
            contexts.append(value.get("contextInfo"))
    for context in contexts:
        if isinstance(context, dict) and context.get("mentionedJid"):
            return True
    return any(re.search(pattern, content or "", re.I) for pattern in BRIDGE_MENTION_PATTERNS)


async def cw_headers() -> dict[str, str]:
    return {"api_access_token": CHATWOOT_TOKEN, "Content-Type": "application/json"}


async def ensure_conversation(jid: str, name: str) -> tuple[str, int]:
    # Evolution can deliver the same event concurrently (especially during
    # history sync). Serialize the contact/conversation creation so both
    # requests cannot create separate Chatwoot conversations.
    async with conversation_lock:
        conn = db()
        found = conn.execute("SELECT source_id, conversation_id FROM contacts WHERE jid=?", (jid,)).fetchone()
        if found:
            conn.close()
            return found[0], int(found[1])
        payload = {"identifier": jid, "name": name or jid, "phone_number": "+" + jid_number(jid), "custom_attributes": {"whatsapp_jid": jid}}
        async with httpx.AsyncClient(timeout=30) as client:
            r = await client.post(f"{CHATWOOT_URL}/public/api/v1/inboxes/{CHATWOOT_INBOX_IDENTIFIER}/contacts", headers=await cw_headers(), json=payload)
            r.raise_for_status()
            contact = r.json()
            source_id = contact.get("contact_inboxes", [{}])[0].get("source_id") or jid
            if not CHATWOOT_ACCOUNT_ID or not CHATWOOT_APP_TOKEN:
                raise RuntimeError("CHATWOOT_ACCOUNT_ID and CHATWOOT_APP_TOKEN are required to create conversations")
            r = await client.post(
                f"{CHATWOOT_URL}/api/v1/accounts/{CHATWOOT_ACCOUNT_ID}/conversations",
                headers={"api_access_token": CHATWOOT_APP_TOKEN, "Content-Type": "application/json"},
                json={"inbox_id": int(CHATWOOT_INBOX_ID), "contact_id": contact["id"], "status": "open"},
            )
            r.raise_for_status()
            conversation_data = r.json()
            internal_id = int(conversation_data["id"])
            conversation_id = int(conversation_data.get("display_id") or internal_id)
        conn.execute("INSERT OR REPLACE INTO contacts (jid, source_id, conversation_id, internal_id) VALUES (?,?,?,?)", (jid, source_id, conversation_id, internal_id))
        conn.commit(); conn.close()
        return source_id, conversation_id


def claim_message(external_id: str) -> bool:
    """Atomically claim an Evolution message before making network calls."""
    conn = db()
    try:
        conn.execute("BEGIN IMMEDIATE")
        inserted = conn.execute(
            "INSERT OR IGNORE INTO messages (external_id, chatwoot_id, created_at) VALUES (?, NULL, ?)",
            (external_id, time.time()),
        ).rowcount
        conn.commit()
        return bool(inserted)
    finally:
        conn.close()


def record_event(external_id: str, event_type: str, jid: str, payload: dict[str, Any], target_id: str = "") -> None:
    """Persist edits, reactions and revocations even when they do not start an agent turn."""
    conn = db()
    conn.execute(
        "INSERT OR REPLACE INTO event_history VALUES (?,?,?,?,?,?)",
        (external_id, event_type, jid, target_id, json.dumps(payload, ensure_ascii=False)[:200000], time.time()),
    )
    conn.commit()
    conn.close()


async def sync_chatwoot_state(event_type: str, jid: str, target_id: str, payload: dict[str, Any]) -> None:
    """Project WhatsApp revoke/read events onto the matching Chatwoot message."""
    if not CHATWOOT_ACCOUNT_ID or not CHATWOOT_APP_TOKEN or not target_id:
        return
    conn = db()
    contact = conn.execute("SELECT conversation_id FROM contacts WHERE jid=?", (jid,)).fetchone()
    message = conn.execute("SELECT chatwoot_id FROM messages WHERE external_id=?", (target_id,)).fetchone()
    conn.close()
    if not contact or not message or not message[0]:
        return
    conversation_id, chatwoot_message_id = int(contact[0]), int(message[0])
    endpoint = f"{CHATWOOT_URL}/api/v1/accounts/{CHATWOOT_ACCOUNT_ID}/conversations/{conversation_id}/messages/{chatwoot_message_id}"
    headers = {"api_access_token": CHATWOOT_APP_TOKEN, "Content-Type": "application/json"}
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            if event_type in {"MESSAGE_DELETE", "MESSAGE_REVOKE"}:
                response = await client.delete(endpoint, headers=headers)
            else:
                raw_status = str(payload.get("status") or payload.get("type") or "").lower()
                status = {"read", "read-self", "read_self", "played"} & {raw_status}
                status = "read" if status else ("delivered" if raw_status in {"delivered", "delivery"} else "sent" if raw_status in {"sent", "server_ack"} else "")
                if event_type in {"RECEIPT", "READ_RECEIPT"} and not status:
                    status = "read"
                if not status:
                    return
                response = await client.patch(endpoint, headers=headers, json={"status": status})
            if response.status_code not in {200, 204}:
                logger.warning("Chatwoot state sync failed event=%s status=%s", event_type, response.status_code)
    except Exception:
        logger.exception("Chatwoot state sync failed event=%s target=%s", event_type, target_id)


def enqueue_outbox(kind: str, number: str, payload: dict[str, Any]) -> None:
    conn = db()
    conn.execute(
        "INSERT INTO outbox (kind,number,payload,next_attempt,created_at) VALUES (?,?,?,?,?)",
        (kind, number, json.dumps(payload, ensure_ascii=False), time.time(), time.time()),
    )
    conn.commit()
    conn.close()


def _outbox_claim_due(limit: int = 10) -> list[tuple[int, str, str, dict[str, Any], int]]:
    conn = db()
    try:
        conn.execute("BEGIN IMMEDIATE")
        rows = conn.execute("SELECT id,kind,number,payload,attempts FROM outbox WHERE status='pending' AND next_attempt<=? ORDER BY id LIMIT ?", (time.time(), limit)).fetchall()
        for row in rows:
            conn.execute("UPDATE outbox SET status='inflight' WHERE id=?", (row[0],))
        conn.commit()
        return [(int(row[0]), row[1], row[2], json.loads(row[3]), int(row[4])) for row in rows]
    finally:
        conn.close()


def _outbox_done(item_id: int) -> None:
    conn = db(); conn.execute("DELETE FROM outbox WHERE id=?", (item_id,)); conn.commit(); conn.close()


def _outbox_failed(item_id: int, attempts: int, error: str) -> None:
    delay = min(900, 2 ** min(attempts + 1, 9))
    conn = db()
    conn.execute("UPDATE outbox SET status='pending', attempts=?, next_attempt=?, last_error=? WHERE id=?", (attempts + 1, time.time() + delay, str(error)[:1000], item_id))
    conn.commit(); conn.close()


async def process_outbox_once() -> int:
    items = await asyncio.to_thread(_outbox_claim_due)
    for item_id, kind, number, payload, attempts in items:
        try:
            if kind == "text":
                await send_evolution(number, str(payload.get("content") or ""), payload.get("reply_to"), retry=True)
            elif kind == "media":
                media = (Path(payload["path"]).read_bytes(), payload["mimetype"], payload["filename"])
                await send_evolution_media(number, media, str(payload.get("caption") or ""), retry=True)
            elif kind.startswith("native:"):
                await send_evolution_native(number, kind.split(":", 1)[1], payload.get("attributes") or {}, retry=True)
            else:
                raise ValueError(f"unknown outbox kind={kind}")
            await asyncio.to_thread(_outbox_done, item_id)
        except Exception as exc:
            logger.warning("Outbox delivery failed id=%s kind=%s attempt=%s: %s", item_id, kind, attempts + 1, exc)
            await asyncio.to_thread(_outbox_failed, item_id, attempts, exc)
    return len(items)


async def outbox_worker() -> None:
    while True:
        try:
            await process_outbox_once()
        except asyncio.CancelledError:
            raise
        except Exception:
            logger.exception("Outbox worker iteration failed")
        await asyncio.sleep(5)


async def add_chatwoot_message(source_id: str, conversation_id: int, content: str, private: bool = False, outgoing: bool = False, media: tuple[bytes, str, str] | None = None, suppress_evolution: bool = False, structured_event: dict[str, Any] | None = None) -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=30) as client:
        if CHATWOOT_ACCOUNT_ID and CHATWOOT_APP_TOKEN:
            endpoint = f"{CHATWOOT_URL}/api/v1/accounts/{CHATWOOT_ACCOUNT_ID}/conversations/{conversation_id}/messages"
            headers = {"api_access_token": CHATWOOT_APP_TOKEN}
            attributes = {"skip_evolution_bridge": True} if suppress_evolution else {}
            if structured_event:
                attributes["hermes_structured_event"] = structured_event
            fields = {"content": content, "message_type": "outgoing" if outgoing else "incoming", "private": "true" if private else "false", "content_attributes": json_dumps(attributes)}
            if media:
                payload, mimetype, filename = media
                r = await client.post(endpoint, headers=headers, data=fields, files={"attachments[]": (filename, payload, mimetype)})
            else:
                r = await client.post(endpoint, headers={**headers, "Content-Type": "application/json"}, json={"content": content, "message_type": "outgoing" if outgoing else "incoming", "private": private, "content_attributes": attributes})
            r.raise_for_status(); return r.json()
        r = await client.post(
            f"{CHATWOOT_URL}/public/api/v1/inboxes/{CHATWOOT_INBOX_IDENTIFIER}/contacts/{source_id}/conversations/{conversation_id}/messages",
            headers=await cw_headers(), json={"content": content, "private": private},
        )
        r.raise_for_status()
        return r.json()


async def ask_hermes(content: str, session: str, media: list[tuple[bytes, str, str]] | None = None, structured_events: list[dict[str, Any]] | None = None, system_context: str = "") -> str:
    if not HERMES_KEY:
        return ""
    headers = {"Authorization": f"Bearer {HERMES_KEY}", "Content-Type": "application/json", "X-Hermes-Session-Id": session}
    parts: list[dict[str, Any]] = [{"type": "text", "text": content}] if content else []
    for payload, mimetype, _filename in media or []:
        # Hermes' OpenAI-compatible endpoint accepts image_url parts (including
        # data URLs), but explicitly rejects uploaded files/audio parts.
        if not mimetype.startswith("image/") or len(payload) > HERMES_MEDIA_MAX_BYTES:
            continue
        encoded = base64.b64encode(payload).decode("ascii")
        parts.append({"type": "image_url", "image_url": {"url": f"data:{mimetype};base64,{encoded}"}})
    message_content: Any = parts if len(parts) > 1 else (content or "[mídia recebida]")
    # The system message is appended to Hermes' prompt for this request only
    # and never persisted; it is the trusted channel for bridge context.
    system = [{"role": "system", "content": system_context}] if system_context else []
    body = {"model": HERMES_MODEL, "messages": [*system, {"role": "user", "content": message_content}], "stream": False}
    if structured_events:
        # Keep a machine-readable copy in the request metadata and a textual
        # copy in the user turn so both Hermes gateway versions can consume it.
        encoded = json.dumps(structured_events, ensure_ascii=False, sort_keys=True)
        marker = f"\n\n[EVENTO_ESTRUTURADO_BRIDGE]\n{encoded}\n[/EVENTO_ESTRUTURADO_BRIDGE]"
        if isinstance(message_content, str):
            message_content += marker
        else:
            message_content = list(message_content) + [{"type": "text", "text": marker}]
        body["metadata"] = {"bridge_structured_events": structured_events}
        body["messages"][-1]["content"] = message_content
    # Fila (07/10): no máximo HERMES_MAX_CONCURRENCY turnos ao mesmo tempo; o resto espera a vez.
    # Se mesmo assim vier 429 (limite de execuções do Hermes), espera e tenta de novo, sem perder a mensagem.
    deadline = time.monotonic() + HERMES_429_RETRY_SECONDS
    delay = 1.0
    async with hermes_slots:
        async with httpx.AsyncClient(timeout=180) as client:
            while True:
                r = await client.post(f"{HERMES_URL}/chat/completions", headers=headers, json=body)
                if getattr(r, "status_code", 200) == 429 and time.monotonic() < deadline:
                    try:
                        wait = float(r.headers.get("Retry-After") or delay)
                    except ValueError:
                        wait = delay
                    logger.warning("Hermes at its concurrent-run cap (429); retrying in %.1fs session=%s", wait, session)
                    await asyncio.sleep(min(max(wait, 0.5), 10.0))
                    delay = min(delay * 2, 10.0)
                    continue
                r.raise_for_status()
                data = r.json()
                break
    return data.get("choices", [{}])[0].get("message", {}).get("content", "")


def extract_hermes_media(reply: str) -> tuple[str, list[tuple[bytes, str, str]]]:
    """Resolve Hermes MEDIA directives through the read-only shared data volume."""
    text = str(reply or "")
    media: list[tuple[bytes, str, str]] = []
    for match in re.finditer(r"MEDIA:\s*(\S+)", text, re.IGNORECASE):
        raw_path = match.group(1)
        candidate = Path(raw_path)
        if raw_path.startswith(HERMES_MEDIA_PREFIX + "/"):
            candidate = HERMES_MEDIA_ROOT / raw_path[len(HERMES_MEDIA_PREFIX) + 1:]
        elif not candidate.is_absolute():
            candidate = HERMES_MEDIA_ROOT / raw_path
        try:
            resolved = candidate.resolve()
            resolved.relative_to(HERMES_MEDIA_ROOT)
            if not resolved.is_file() or resolved.stat().st_size > HERMES_MEDIA_MAX_BYTES:
                continue
            payload = resolved.read_bytes()
            mimetype = guess_type(resolved.name)[0] or "application/octet-stream"
            media.append((payload, mimetype, resolved.name))
        except (OSError, ValueError):
            logger.warning("Ignoring Hermes media directive outside shared root: %s", raw_path)
    cleaned = re.sub(r"MEDIA:\s*\S+", "", text, flags=re.IGNORECASE)
    return cleaned.strip(), media


async def send_presence(number: str, state: str = "composing") -> None:
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            await client.post(f"{EVOLUTION_URL}/message/presence", headers={"apikey": EVOLUTION_API_KEY}, json={"number": number, "state": state, "isAudio": False})
    except Exception:
        logger.debug("Presence update failed for %s", number, exc_info=True)


async def escalate(jid: str, source_id: str, conversation_id: int, directive: dict[str, Any], last_message: str) -> None:
    """Hand the conversation to the owner: note, labels, pause and a WhatsApp alert.

    Idempotent: the agent's [[ESCALAR]] and the supervisor can both ask for it;
    once the conversation is paused the owner already has it.
    """
    try:
        if await handover.paused(conversation_id):
            logger.info("SDR escalation skipped, conversation=%s already paused (motivo=%s)", conversation_id, directive.get("motivo"))
            return
    except Exception:
        logger.warning("Could not read pause state for conversation=%s; escalating anyway", conversation_id)
    logger.warning("SDR escalation conversation=%s motivo=%s", conversation_id, directive.get("motivo"))
    for step in (
        lambda: sdr.private_note(conversation_id, sdr.escalation_note(directive, last_message)),
        lambda: sdr.add_label(conversation_id, sdr.ESCALATION_LABEL),
        lambda: handover.pause(conversation_id),
    ):
        try:
            await step()
        except Exception:
            logger.exception("Escalation step failed for conversation=%s", conversation_id)
    if sdr.NOTIFY_NUMBER:
        try:
            await send_evolution(sdr.NOTIFY_NUMBER, sdr.escalation_alert(directive, jid_number(jid), conversation_id))
        except Exception:
            logger.exception("Owner escalation alert failed for conversation=%s", conversation_id)

_background_tasks: set[asyncio.Task] = set()


def fire_and_forget(coro: Any) -> None:
    """Run a coroutine without awaiting it, keeping a reference until it ends."""
    task = asyncio.create_task(coro)
    _background_tasks.add(task)
    task.add_done_callback(_background_tasks.discard)


async def brain_event(path: str, payload: dict[str, Any]) -> None:
    if not SDR_BRAIN_URL: return
    try:
        async with httpx.AsyncClient(timeout=5) as client:
            await client.post(SDR_BRAIN_URL + path, headers={"Authorization": f"Bearer {SDR_BRAIN_TOKEN}"}, json=payload)
    except Exception: logger.debug("SDR event failed", exc_info=True)

async def record_offer_from_origin(conversation_id: int, origin: dict[str, Any] | None) -> None:
    """The campaign text names the offer: record it like a [[REGISTRO: oferta = ...]]."""
    offer = (origin or {}).get("oferta")
    if offer not in sdr.OFFERS:
        return
    await sdr.record(conversation_id, "oferta", offer)
    fire_and_forget(brain_event("/events/record", {"conversation_id": conversation_id, "campo": "oferta", "valor": offer, "ts": time.time()}))


async def review_reply(jid: str, conversation_id: int, last_message: str, parts: list[str], directives: list[dict[str, Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    if not SDR_BRAIN_URL: return parts, directives
    try:
        async with httpx.AsyncClient(timeout=SDR_REVIEW_TIMEOUT) as client:
            r = await client.post(SDR_BRAIN_URL + "/review", headers={"Authorization": f"Bearer {SDR_BRAIN_TOKEN}"}, json={"conversation_id":conversation_id,"jid":jid,"lead_message":last_message,"draft_parts":parts,"directives":directives,"context":sdr.context_payload(conversation_id)})
            data = r.json(); decision = data.get("decision")
            if decision == "reescrever":
                rewritten = data.get("parts") or parts
                np, nd = sdr.parse_reply("[[QUEBRA]]".join(map(str, rewritten)))
                return np, (nd or directives)
            if decision == "bloquear":
                await sdr.private_note(conversation_id, "Resposta bloqueada pelo supervisor: " + str(data.get("motivo") or "não informado"))
                return [sdr.FALLBACK_REPLY], [{"type":"escalar","motivo":str(data.get("motivo") or "revisão bloqueou"),"frase":last_message}]
    except Exception: logger.warning("SDR review failed", exc_info=True)
    return parts, directives


async def send_credential_notice(jid: str, source_id: str, conversation_id: int) -> bool:
    """The standard notice for a pasted credential, sent by the bridge itself.

    The secret was already replaced in everything the agent, the brain and
    Chatwoot see; WhatsApp keeps the lead's copy, so he is asked to delete it.
    """
    try:
        await send_evolution(jid_number(jid), sdr.CREDENTIAL_NOTICE)
        await add_chatwoot_message(source_id, conversation_id, sdr.CREDENTIAL_NOTICE, private=False, outgoing=True, suppress_evolution=True)
        sdr.touch_activity(conversation_id, agent=True)
        await sdr.record(conversation_id, "mandou_credencial", "sim")
        fire_and_forget(brain_event("/events/record", {"conversation_id": conversation_id, "campo": "mandou_credencial", "valor": "sim", "ts": time.time()}))
        return True
    except Exception:
        logger.exception("Credential notice failed for conversation=%s", conversation_id)
        return False


# Before the purchase, a lead with the profile for the Raio-X is recorded and
# the sale goes on; only after it (or when he asks) the Talles is called.
# Only the agent's own "[[ESCALAR: perfil para Raio-X | …]]": a block by the
# reviewer may cite the Raio-X in its motivo and must still reach a person.
PROFILE_ESCALATION_RE = re.compile(r"^\s*perfil\b", re.I)
ASKED_FOR_TALLES_RE = re.compile(r"raio[\s-]?x|copiar a minha opera|humano|\bfalar com (?:o talles|uma pessoa|algu[eé]m|voc[eê]s)|atendente", re.I)
ANNOUNCES_TALLES_RE = re.compile(r"\b(?:vou chamar|chamo|vou passar|passo)\b[^.!?\n]{0,30}\b(?:talles|ele)\b", re.I)


def hold_profile_escalation(conversation_id: int, parts: list[str], directives: list[dict[str, Any]], last_message: str) -> tuple[list[str], list[dict[str, Any]]]:
    escalation = next((d for d in directives if d["type"] == "escalar"), None)
    if not escalation or not PROFILE_ESCALATION_RE.search(escalation.get("motivo", "")) or ASKED_FOR_TALLES_RE.search(last_message or ""):
        return parts, directives
    state = sdr.lead_state(conversation_id)
    bought = sdr._fold(str((state.get("registros") or {}).get("comprou") or "")).strip() == "sim"
    if state.get("oferta") != "instalador" or bought:
        return parts, directives
    kept = [part for part in parts if not ANNOUNCES_TALLES_RE.search(part)]
    if not kept:
        return parts, directives  # nothing else to say: a person takes it, as the agent announced
    logger.info("SDR profile escalation held until the purchase conversation=%s", conversation_id)
    return kept, [d for d in directives if d is not escalation] + [{"type": "registro", "campo": "perfil_raiox", "valor": "sim"}]


def lead_interrupted(jid: str) -> bool:
    """The lead wrote again while this reply was being prepared or sent."""
    return bool(hermes_pending.get(jid))


async def deliver_agent_reply(jid: str, source_id: str, conversation_id: int, reply: str, reply_to_id: str | None, last_message: str) -> bool:
    """Parse directives, apply guardrails, then send like a person typing.

    Returns False when the lead interrupted before anything was sent: the
    caller then answers the old and the new messages together instead of
    replying late to the old ones.
    """
    reply_text, reply_media = extract_hermes_media(reply)
    parts, directives = sdr.parse_reply(reply_text)
    draft = list(parts)
    if parts:
        parts, directives = await review_reply(jid, conversation_id, last_message, parts, directives)
    parts, directives = hold_profile_escalation(conversation_id, parts, directives, last_message)
    leaks = sdr.leaked_terms(parts)
    if leaks:
        logger.warning("SDR guardrail blocked reply conversation=%s terms=%s", conversation_id, leaks)
        parts = [sdr.FALLBACK_REPLY]
        directives = [d for d in directives if d["type"] == "registro"] + [{"type": "escalar", "motivo": f"guardrail bloqueou termos internos: {', '.join(leaks)}", "frase": last_message}]
        reply_media = []
        try:
            await sdr.private_note(conversation_id, "🛑 Resposta do agente bloqueada pelo guardrail (termos: " + ", ".join(leaks) + "):\n\n" + reply_text)
        except Exception:
            logger.exception("Failed to store blocked reply for conversation=%s", conversation_id)
    # Deterministic style rule from the owner: no dashes or colons chaining sentences.
    parts = [sdr.plain_style(part) for part in parts]
    number = jid_number(jid)
    sent: list[str] = []
    for index, part in enumerate(parts):
        if lead_interrupted(jid):
            break
        await send_presence(number)
        await asyncio.sleep(sdr.typing_seconds(part))
        if lead_interrupted(jid):
            break
        await send_evolution(number, part, reply_to=reply_to_id if index == 0 else None)
        await add_chatwoot_message(source_id, conversation_id, part, private=False, outgoing=True, suppress_evolution=True)
        sdr.touch_activity(conversation_id, agent=True)
        sent.append(part)
    # Hermes keeps its draft in the session; tell it what the lead really got.
    sdr.record_correction(conversation_id, draft, sent, "interrompida pelo lead" if len(sent) < len(parts) else "review_or_guardrail")
    if len(sent) < len(parts):
        logger.info("SDR reply interrupted by the lead conversation=%s sent=%d/%d", conversation_id, len(sent), len(parts))
        await send_presence(number, "paused")
        if not sent:
            return False
        # Part of the answer went out: keep the facts it established, skip the
        # rest (media, escalation) and let the next turn continue from there.
        for directive in directives:
            if directive["type"] == "registro":
                await sdr.record(conversation_id, directive["campo"], directive["valor"])
        return True
    photos = [sdr.load_media(d["nome"]) for d in directives if d["type"] == "foto"]
    for attachment in [*reply_media, *[p for p in photos if p]]:
        await send_evolution_media(number, attachment, retry=False)
        await add_chatwoot_message(source_id, conversation_id, "", private=False, outgoing=True, media=attachment, suppress_evolution=True)
    for directive in directives:
        if directive["type"] == "registro":
            await sdr.record(conversation_id, directive["campo"], directive["valor"])
            fire_and_forget(brain_event("/events/record", {"conversation_id":conversation_id,"campo":directive["campo"],"valor":directive["valor"],"ts":time.time()}))
    escalation = next((d for d in directives if d["type"] == "escalar"), None)
    if escalation:
        await escalate(jid, source_id, conversation_id, escalation, last_message)
    fire_and_forget(brain_event("/events/exchange", {"conversation_id":conversation_id,"jid":jid,"source_id":source_id,"lead_message":last_message,"agent_parts":parts,"directives":directives,"ts":time.time()}))
    return True


async def hermes_batch_worker(jid: str) -> None:
    """Coalesce rapid messages from one JID into one Hermes turn.

    Chatwoot receives each inbound message immediately; only the agent turn is
    delayed. Hermes then persists the combined turn under the stable wa:<jid>
    session, so later turns retain the conversation history.
    """
    try:
        while True:
            await asyncio.sleep(HERMES_MESSAGE_DEBOUNCE_SECONDS)
            batch = hermes_pending.pop(jid, [])
            if batch:
                texts = [text for text, _media, _source, _conversation, _external_id, _event in batch if text and text != "[mídia recebida]"]
                combined = "\n".join(texts)
                structured_events = [event for *_rest, event in batch if event]
                media = [attachment for _text, items, _source, _conversation, _external_id, _event in batch for attachment in (items or [])]
                transcripts = await transcribe_audio_media(media)
                if transcripts:
                    texts.extend(f"[Transcrição do áudio]: {text}" for text in transcripts)
                    combined = "\n".join(texts)
                document_text = await asyncio.to_thread(extract_document_media, media)
                if document_text:
                    texts.extend(document_text)
                    combined = "\n".join(texts)
                source_id, conversation_id = batch[-1][2], batch[-1][3]
                reply_to_id = batch[-1][4]
                if not media and any(re.search(r"\b(imagem|foto|figurinha|sticker|áudio|audio|vídeo|video|arquivo|documento|mídia|media|isso|essa)\b", text, re.I) for text in texts):
                    media = recent_local_media(jid)
                    if not media:
                        media = await recent_chatwoot_media(conversation_id)
                if not combined and media:
                    combined = "[mídia recebida]"
                # Also covers a password said in an audio or written in a document.
                combined, secret = sdr.redact_credentials(combined)
                logger.info("Hermes batch jid=%s messages=%d media=%d structured_events=%d history_media=%s", jid, len(batch), len(media), len(structured_events), bool(media and not any(item for _text, item, _source, _conversation, _external_id, _event in batch)))
                try:
                    if await handover.paused(conversation_id):
                        handover.record(conversation_id, uuid.uuid4().hex, "cliente", combined)
                        continue
                    version = handover.revision(conversation_id)
                    history, through = handover.context(conversation_id)
                    prompt = (history + "\nMensagem atual do cliente:\n" + combined) if history else combined
                    first = sdr.remember_lead(conversation_id, jid, combined)
                    sdr.touch_activity(conversation_id, lead=True)
                    sdr.cancel_followups(conversation_id, "lead_message")
                    if first:
                        origin = sdr.lead_state(conversation_id).get("origem")
                        await record_offer_from_origin(conversation_id, origin)
                        fire_and_forget(brain_event("/events/lead", {"conversation_id":conversation_id,"jid":jid,"origin":origin,"first_message":combined,"ts":time.time()}))
                    notice_sent = await send_credential_notice(jid, source_id, conversation_id) if secret else False
                    guidance = sdr.pending_guidance(conversation_id)
                    reply = await ask_hermes(prompt, HERMES_SESSION_PREFIX + jid + sdr.session_suffix(jid), media, structured_events, sdr.trusted_context(conversation_id, guidance, credential_notice_sent=notice_sent))
                    if await handover.paused(conversation_id) or version != handover.revision(conversation_id):
                        handover.record(conversation_id, uuid.uuid4().hex, "cliente", combined)
                        continue
                    if reply and lead_interrupted(jid):
                        # The lead wrote while Hermes was thinking: this reply
                        # answers an outdated question. Drop it and answer the
                        # old and the new messages together on the next pass.
                        parts_draft, _ = sdr.parse_reply(extract_hermes_media(reply)[0])
                        sdr.record_correction(conversation_id, parts_draft, [], "interrompida pelo lead")
                        hermes_pending[jid][:0] = batch
                        logger.info("SDR reply discarded, lead interrupted conversation=%s", conversation_id)
                        continue
                    if reply:
                        delivered = await deliver_agent_reply(jid, source_id, conversation_id, reply, reply_to_id, combined)
                        if not delivered:
                            hermes_pending[jid][:0] = batch
                            continue
                        sdr.consume_guidance([gid for gid, _text in guidance])
                        handover.consumed(conversation_id, through)
                except Exception:
                    handover.record(conversation_id, uuid.uuid4().hex, "cliente", combined)
                    logger.exception("Hermes/Evolution batch reply failed for jid=%s", jid)
            if not hermes_pending.get(jid):
                return
    finally:
        hermes_tasks.pop(jid, None)


async def owner_test_command(jid: str, command: str) -> None:
    """#teste starts a fresh lead conversation for the owner; #fim ends it."""
    conn = db()
    row = conn.execute("SELECT conversation_id FROM contacts WHERE jid=?", (jid,)).fetchone()
    if row:
        sdr.cancel_followups(int(row[0]), "owner_test_" + command.lstrip("#"))
    if command == sdr.OWNER_TEST_START:
        run = sdr.start_owner_test(jid)
        # Dropping the mapping makes the next message open a new Chatwoot
        # conversation, so labels, records and transcript start clean.
        conn.execute("DELETE FROM contacts WHERE jid=?", (jid,))
        conn.commit()
        text = (f"🧪 Modo teste {run} ligado. A partir da próxima mensagem você fala com o agente como um lead novo, "
                "numa conversa nova. Mande #fim para sair.")
    else:
        sdr.stop_owner_test(jid)
        text = "🧪 Modo teste desligado. Seu número volta a ser tratado como dono: mensagens daqui não vão ao agente."
    conn.close()
    hermes_pending.pop(jid, None)
    logger.info("Owner test command %s for %s", command, jid)
    try:
        await send_evolution(sdr.NOTIFY_NUMBER or jid_number(jid), text)
    except Exception:
        logger.exception("Owner test confirmation failed")


def queue_hermes_turn(jid: str, content: str, media: list[tuple[bytes, str, str]] | None, source_id: str, conversation_id: int, external_id: str, structured_event: dict[str, Any] | None = None) -> None:
    sdr.touch_activity(conversation_id, lead=True)
    sdr.cancel_followups(conversation_id, "lead_message")
    hermes_pending[jid].append((content, media, source_id, conversation_id, external_id, structured_event))
    if jid not in hermes_tasks:
        hermes_tasks[jid] = asyncio.create_task(hermes_batch_worker(jid))

async def process_followups_once(now: float | None = None) -> None:
    """Send due follow-ups written by sdr-brain, applying deterministic guardrails."""
    if not SDR_BRAIN_URL: return
    from zoneinfo import ZoneInfo
    from datetime import datetime
    now = time.time() if now is None else now
    tz = ZoneInfo("America/Sao_Paulo")
    local = datetime.fromtimestamp(now, tz)
    day_start = local.replace(hour=0, minute=0, second=0, microsecond=0).timestamp()
    for row in sdr.pending_followups():
        fid, conv, send_at, text, motivo, after_ts, stage, *_ = row
        if send_at > now: continue
        with sdr.connect() as c:
            still_pending = c.execute("SELECT 1 FROM sdr_followups WHERE id=? AND status='pending'", (fid,)).fetchone()
        if not still_pending:
            continue  # superseded earlier in this pass
        conn = db(); contact = conn.execute("SELECT jid,source_id FROM contacts WHERE conversation_id=?", (conv,)).fetchone(); conn.close()
        if not contact or (sdr.is_owner(contact[0]) and not sdr.owner_test_run(contact[0])): sdr.update_followup(fid,"cancelled","owner_or_missing"); continue
        last_lead, _ = sdr.activity(conv)
        if last_lead and last_lead > after_ts: sdr.update_followup(fid,"cancelled","lead_message"); continue
        try:
            if await handover.paused(conv): sdr.update_followup(fid,"cancelled","paused"); continue
        except Exception: continue
        with sdr.connect() as c:
            if c.execute("SELECT 1 FROM sdr_records WHERE conversation=? AND field='desfecho'",(conv,)).fetchone(): sdr.update_followup(fid,"cancelled","desfecho"); continue
            touches = c.execute("SELECT COUNT(*) FROM sdr_followups WHERE conversation=? AND status='sent' AND sent_at>=?",(conv, day_start)).fetchone()[0]
        if touches >= 4: sdr.update_followup(fid,"cancelled","daily_cap"); continue
        if local.hour < SDR_BUSINESS_START_HOUR or local.hour >= SDR_BUSINESS_END_HOUR: continue
        leaks = sdr.leaked_terms([text])
        if leaks:
            sdr.update_followup(fid, "cancelled", "guardrail:" + ",".join(leaks))
            try:
                await sdr.private_note(conv, "🛑 Retomada bloqueada pelo guardrail (termos: " + ", ".join(leaks) + "):\n\n" + text)
            except Exception:
                logger.exception("Failed to store blocked follow-up for conversation=%s", conv)
            continue
        text = sdr.plain_style(text)
        try:
            number = jid_number(contact[0])
            await send_presence(number, "composing")
            await asyncio.sleep(sdr.typing_seconds(text))
            await send_evolution(number, text)
            await add_chatwoot_message(contact[1], conv, text, outgoing=True, suppress_evolution=True)
            # Visible to the SDR on the next turn. The silence cadence stays
            # anchored on the last conversational agent reply (last_agent_ts),
            # so a touch never restarts the cadence.
            handover.record(conv, "followup:" + str(fid), "retomada_enviada", text)
            sdr.update_followup(fid, "sent")
            # One touch per silence window: anything else still queued for this
            # conversation (e.g. decided overnight) is superseded.
            sdr.cancel_followups(conv, "superseded")
        except Exception:
            logger.exception("Follow-up %s failed for conversation=%s", fid, conv)

async def followup_worker() -> None:
    while True:
        try: await process_followups_once()
        except asyncio.CancelledError: raise
        except Exception: logger.exception("followup iteration failed")
        await asyncio.sleep(30)


def split_outbound_text(content: str, limit: int = EVOLUTION_MAX_MESSAGE_LENGTH) -> list[str]:
    """Split WhatsApp text at readable boundaries while respecting Evolution's limit."""
    text = str(content or "")
    if len(text) <= limit:
        return [text]
    chunks = []
    while text:
        if len(text) <= limit:
            chunks.append(text)
            break
        cut = max(text.rfind("\n", 0, limit + 1), text.rfind(" ", 0, limit + 1))
        cut = cut if cut > limit // 2 else limit
        chunks.append(text[:cut].rstrip())
        text = text[cut:].lstrip()
    return [chunk for chunk in chunks if chunk]


def format_whatsapp_text(content: str) -> str:
    """Translate Hermes Markdown to the formatting accepted by WhatsApp."""
    text = str(content or "")
    saved: list[str] = []

    def stash(match: re.Match[str]) -> str:
        saved.append(match.group(0))
        return f"\x00CODE{len(saved) - 1}\x00"

    text = re.sub(r"```[\s\S]*?```", stash, text)
    text = re.sub(r"`[^`\n]+`", stash, text)
    text = re.sub(r"(?<!\*)\*(?!\s|\*)([^*\n]*?\S[^*\n]*?)\*(?!\*)", r"_\1_", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"*\1*", text, flags=re.DOTALL)
    text = re.sub(r"__(.+?)__", r"*\1*", text, flags=re.DOTALL)
    text = re.sub(r"~~(.+?)~~", r"~\1~", text, flags=re.DOTALL)
    text = re.sub(r"^#{1,6}\s+(.+)$", r"*\1*", text, flags=re.MULTILINE)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for index, original in enumerate(saved):
        text = text.replace(f"\x00CODE{index}\x00", original)
    return text


async def send_evolution(number: str, content: str, reply_to: str | None = None, retry: bool = False) -> None:
    send_path = EVOLUTION_SEND_PATH.replace("{instance}", EVOLUTION_INSTANCE)
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            chunks = split_outbound_text(format_whatsapp_text(content))
            for index, chunk in enumerate(chunks):
                outbound_id = uuid.uuid4().hex.upper()
                handover.outbound_id(outbound_id)
                payload = {"number": number, "text": chunk, "instance": EVOLUTION_INSTANCE, "id": outbound_id}
                if reply_to and index == 0:
                    payload["quotedMessageId"] = reply_to
                    payload["replyTo"] = reply_to
                r = await client.post(f"{EVOLUTION_URL}{send_path}", headers={"apikey": EVOLUTION_API_KEY}, json=payload)
                r.raise_for_status()
                if len(chunks) > 1:
                    await asyncio.sleep(0.3)
    except Exception:
        if not retry:
            enqueue_outbox("text", number, {"content": content, "reply_to": reply_to})
        raise


async def send_evolution_native(number: str, kind: str, payload: dict[str, Any], retry: bool = False) -> None:
    """Send a native Evolution primitive with modern and classic endpoint fallbacks."""
    paths = [f"/send/{kind}".replace("{instance}", EVOLUTION_INSTANCE), f"/message/send{kind.title()}/{EVOLUTION_INSTANCE}"]
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            for index, path in enumerate(dict.fromkeys(paths)):
                response = await client.post(f"{EVOLUTION_URL}{path}", headers={"apikey": EVOLUTION_API_KEY}, json={**payload, "number": number, "instance": EVOLUTION_INSTANCE})
                if response.status_code == 404 and index + 1 < len(paths):
                    continue
                response.raise_for_status()
                return
        raise RuntimeError(f"Evolution native send failed kind={kind}")
    except Exception:
        if not retry:
            enqueue_outbox(f"native:{kind}", number, {"attributes": payload})
        raise


async def send_evolution_location(number: str, attributes: dict[str, Any], retry: bool = False) -> None:
    await send_evolution_native(number, "location", {
        "latitude": float(attributes["latitude"]), "longitude": float(attributes["longitude"]),
        "name": attributes.get("name", ""), "address": attributes.get("address", ""),
    }, retry=retry)


async def send_evolution_poll(number: str, attributes: dict[str, Any], retry: bool = False) -> None:
    options = [str(item).strip() for item in attributes.get("options", []) if str(item).strip()]
    if not attributes.get("question") or len(options) < 2:
        raise ValueError("Poll requires question and at least two options")
    await send_evolution_native(number, "poll", {
        "question": str(attributes["question"]), "options": options,
        "selectableCount": int(attributes.get("selectable_count", attributes.get("selectableCount", 1))),
    }, retry=retry)


async def send_evolution_reaction(number: str, attributes: dict[str, Any], retry: bool = False) -> None:
    message_id = str(attributes.get("message_id") or attributes.get("messageId") or "")
    if not message_id:
        raise ValueError("Reaction requires message_id")
    await send_evolution_native(number, "reaction", {
        "key": {"remoteJid": f"{number}@s.whatsapp.net", "id": message_id},
        "reaction": str(attributes.get("emoji") or attributes.get("reaction") or ""),
    }, retry=retry)


def _attachment_mimetype(attachment: dict[str, Any]) -> str:
    value = attachment.get("mime_type") or attachment.get("content_type") or attachment.get("mimetype")
    if value and "/" in str(value):
        return str(value)
    file_type = str(attachment.get("file_type") or "").lower()
    return {"image": "image/jpeg", "video": "video/mp4", "audio": "audio/ogg", "file": "application/octet-stream"}.get(file_type, "application/octet-stream")


async def download_chatwoot_attachment(attachment: dict[str, Any]) -> tuple[bytes, str, str] | None:
    """Fetch a Chatwoot attachment and normalize it for Evolution's multipart API."""
    if not isinstance(attachment, dict):
        return None
    url = next((attachment.get(key) for key in ("data_url", "file_url", "url") if attachment.get(key)), "")
    if not url:
        return None
    filename = str(attachment.get("file_name") or attachment.get("filename") or "attachment")
    mimetype = _attachment_mimetype(attachment)
    if url.startswith("data:") and "," in url:
        header, encoded = url.split(",", 1)
        mimetype = header[5:].split(";", 1)[0] or mimetype
        try:
            payload = base64.b64decode(encoded)
        except (ValueError, base64.binascii.Error):
            return None
    else:
        headers = {"api_access_token": CHATWOOT_APP_TOKEN} if CHATWOOT_APP_TOKEN else {}
        async with httpx.AsyncClient(timeout=30, follow_redirects=True) as client:
            response = await client.get(str(url), headers=headers)
            response.raise_for_status()
            payload = response.content
            mimetype = response.headers.get("content-type", mimetype).split(";", 1)[0]
    if not payload:
        return None
    if "." not in filename.rsplit("/", 1)[-1]:
        filename += guess_extension(mimetype) or ".bin"
    return payload, mimetype, filename


async def recent_chatwoot_media(conversation_id: int, limit: int = 3) -> list[tuple[bytes, str, str]]:
    """Recover recent attachments from Chatwoot ActiveStorage for a media reference."""
    if not CHATWOOT_ACCOUNT_ID or not CHATWOOT_APP_TOKEN:
        return []
    headers = {"api_access_token": CHATWOOT_APP_TOKEN}
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.get(f"{CHATWOOT_URL}/api/v1/accounts/{CHATWOOT_ACCOUNT_ID}/conversations/{conversation_id}/messages", headers=headers)
            response.raise_for_status()
            messages = response.json().get("payload", [])
    except Exception:
        logger.exception("Chatwoot media history lookup failed for conversation=%s", conversation_id)
        return []
    result: list[tuple[bytes, str, str]] = []
    for message in reversed(messages):
        for attachment in message.get("attachments") or []:
            media = await download_chatwoot_attachment(attachment)
            if media:
                result.append(media)
                if len(result) >= limit:
                    return result
    return result


async def send_evolution_media(number: str, media: tuple[bytes, str, str], caption: str = "", retry: bool = False) -> None:
    payload, mimetype, filename = media
    kind = mimetype.split("/", 1)[0]
    if kind not in {"image", "video", "audio"}:
        kind = "document"
    send_path = EVOLUTION_SEND_MEDIA_PATH.replace("{instance}", EVOLUTION_INSTANCE)
    outbound_id = uuid.uuid4().hex.upper()
    handover.outbound_id(outbound_id)
    fields = {
        "id": outbound_id,
        "number": number,
        "type": kind,
        "caption": caption,
        "filename": filename,
        "instance": EVOLUTION_INSTANCE,
    }
    paths = [send_path]
    classic_path = f"/message/sendMedia/{EVOLUTION_INSTANCE}"
    if classic_path not in paths:
        paths.append(classic_path)
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            for index, path in enumerate(paths):
                response = await client.post(
                    f"{EVOLUTION_URL}{path}",
                    headers={"apikey": EVOLUTION_API_KEY},
                    data=fields,
                    files={"file": (filename, payload, mimetype)},
                )
                if response.status_code == 404 and index + 1 < len(paths):
                    continue
                response.raise_for_status()
                # Keep outbound attachments addressable when the user replies to them.
                remember_media(f"{number}@s.whatsapp.net", outbound_id, media)
                return
    except Exception:
        if not retry:
            outbound_path = MEDIA_DIR / f"outbox-{outbound_id}{Path(filename).suffix or '.bin'}"
            MEDIA_DIR.mkdir(parents=True, exist_ok=True)
            outbound_path.write_bytes(payload)
            enqueue_outbox("media", number, {"path": str(outbound_path), "mimetype": mimetype, "filename": filename, "caption": caption})
        raise


@app.get("/health")
async def health() -> dict[str, Any]:
    conn = db()
    pending = conn.execute("SELECT COUNT(*) FROM outbox WHERE status IN ('pending','inflight')").fetchone()[0]
    conn.close()
    return {"status": "ok", "outbox_pending": int(pending), "stt": STT_ENABLED, "ocr": OCR_ENABLED}

def _sdr_auth(request: Request):
    if not SDR_BRAIN_TOKEN: return JSONResponse({"detail":"SDR brain unavailable"}, status_code=503)
    if request.headers.get("authorization") != "Bearer " + SDR_BRAIN_TOKEN: return JSONResponse({"detail":"unauthorized"}, status_code=401)
    return None

@app.post("/sdr/guidance")
async def sdr_guidance(request: Request):
    if (e:=_sdr_auth(request)): return e
    p=await request.json();
    with sdr.connect() as c: c.execute("INSERT INTO sdr_guidance(conversation,text,created) VALUES (?,?,?)", (int(p["conversation_id"]),str(p["text"]),time.time()))
    return {"ok":True}

@app.post("/sdr/escalate")
async def sdr_escalate(request: Request):
    if (e:=_sdr_auth(request)): return e
    p=await request.json(); conv=int(p["conversation_id"]); conn=db(); row=conn.execute("SELECT jid,source_id FROM contacts WHERE conversation_id=?",(conv,)).fetchone(); conn.close()
    if not row: return JSONResponse({"detail":"conversation not found"},404)
    await escalate(row[0],row[1],conv,{"type":"escalar","motivo":p.get("motivo",""),"frase":p.get("frase","")},p.get("frase","")); return {"ok":True}

@app.post("/sdr/lead-profile")
async def sdr_profile(request: Request):
    if (e:=_sdr_auth(request)): return e
    p=await request.json(); sdr.set_profile(int(p["conversation_id"]),p.get("profile") or {}); return {"ok":True}

@app.post("/sdr/followups")
async def sdr_followup(request: Request):
    if (e:=_sdr_auth(request)): return e
    p=await request.json(); fid=sdr.add_followup(int(p["conversation_id"]),float(p["send_at"]),str(p["text"]),str(p.get("motivo", "")),float(p.get("after_ts",0)),str(p.get("stage", ""))); return {"id":fid}

@app.delete("/sdr/followups/{conversation_id}")
async def sdr_followup_delete(conversation_id: int, request: Request):
    if (e:=_sdr_auth(request)): return e
    return {"cancelled":sdr.cancel_followups(conversation_id,"api")}

@app.get("/sdr/conversations/{conversation_id}")
async def sdr_conversation(conversation_id: int, request: Request):
    if (e:=_sdr_auth(request)): return e
    conn=db(); row=conn.execute("SELECT jid,source_id FROM contacts WHERE conversation_id=?",(conversation_id,)).fetchone(); conn.close()
    if not row: return JSONResponse({"detail":"conversation not found"},404)
    lead,agent=sdr.activity(conversation_id)
    f=[{"id":r[0],"send_at":r[2],"text":r[3],"motivo":r[4],"after_ts":r[5],"stage":r[6],"status":r[7],"reason":r[8],"sent_at":r[9]} for r in sdr.pending_followups(conversation_id)]
    return {"jid":row[0],"source_id":row[1],"lead":sdr.lead_state(conversation_id),"followups":f,"paused":await handover.paused(conversation_id),"last_lead_ts":lead,"last_agent_ts":agent,"context":sdr.context_payload(conversation_id)}


@app.post("/evolution/webhook")
async def evolution_webhook(request: Request) -> dict[str, bool]:
    if BRIDGE_WEBHOOK_TOKEN and request.headers.get("authorization") != f"Bearer {BRIDGE_WEBHOOK_TOKEN}":
        return {"received": False}
    payload = await request.json()
    event = str(payload.get("event", "")).upper()
    data = payload.get("data") or {}
    info = data.get("Info") or data.get("info") or {}
    if event == "BUTTONCLICK":
        raw_key = {"id": data.get("messageId") or data.get("id"), "remoteJid": data.get("jid") or data.get("chat"), "fromMe": data.get("fromMe", False)}
        message = {"templateButtonReplyMessage": {"selectedDisplayText": data.get("buttonText"), "selectedId": data.get("buttonId")}}
    elif event in {"REACTION", "MESSAGE_REACTION"}:
        raw_key = data.get("key") or {"id": data.get("id") or data.get("messageId"), "remoteJid": data.get("jid") or data.get("chat"), "fromMe": data.get("fromMe", False)}
        message = {"reactionMessage": data.get("reactionMessage") or data.get("reaction") or data}
    else:
        raw_key = data.get("key") or {}
        message = data.get("message") or data.get("Message") or {}
    key = raw_key or {
        "id": info.get("ID") or info.get("Id") or data.get("id"),
        "remoteJid": info.get("Chat") or info.get("RemoteJid") or info.get("SenderAlt"),
        "fromMe": info.get("IsFromMe", False),
    }
    if not message:
        message = data.get("message") or data.get("Message") or {}
    push_name = data.get("pushName") or data.get("PushName") or info.get("PushName") or ""
    logger.warning(
        "Evolution webhook event=%s data_type=%s data_keys=%s message_keys=%s",
        event,
        type(data).__name__,
        list(data.keys()) if isinstance(data, dict) else None,
        list(message.keys()) if isinstance(message, dict) else None,
    )
    external_id = str(key.get("id") or data.get("id") or "")
    if event in {"RECEIPT", "READ_RECEIPT", "MESSAGE_STATUS"}:
        receipt_ids = data.get("MessageIDs") or data.get("messageIds") or [external_id]
        receipt_jid = str(data.get("Chat") or key.get("remoteJid") or "")
        for receipt_id in receipt_ids:
            if receipt_id:
                record_event(str(receipt_id), event, receipt_jid, data)
                await sync_chatwoot_state(event, receipt_jid, str(receipt_id), data)
        return {"received": True}
    remote_jid = str(key.get("remoteJid") or "")
    if event not in {"MESSAGE", "SENDMESSAGE", "SEND_MESSAGE", "MESSAGES_UPSERT", "BUTTONCLICK", "REACTION", "MESSAGE_REACTION", "MESSAGE_EDIT", "MESSAGE_DELETE", "MESSAGE_REVOKE", "RECEIPT", "READ_RECEIPT", "MESSAGE_STATUS"} or not external_id or not group_is_allowed(remote_jid):
        return {"received": True}
    from_me = bool(key.get("fromMe"))
    # SENDMESSAGE is Evolution API delivery; MESSAGE/fromMe also includes
    # linked-device sends. IDs registered before API calls suppress only echoes.
    if event in {"SENDMESSAGE", "SEND_MESSAGE"} or handover.is_echo(external_id):
        return {"received": True}
    jid = re.sub(r":\d+(?=@)", "", str(key.get("remoteJid", "")))
    if not jid or jid.endswith(("@broadcast", "@newsletter")) or not group_is_allowed(jid):
        return {"received": True}
    if jid.endswith("@g.us") and not group_triggered(message, text_from_message(message)):
        logger.info("Ignoring unmentioned group message jid=%s", jid)
        return {"received": True}
    alternate = info.get("RecipientAlt") if from_me else info.get("SenderAlt")
    if jid.endswith("@lid") and alternate:
        jid = re.sub(r":\d+(?=@)", "", str(alternate))
    if sdr.is_owner(jid):
        # The owner's own chat carries escalation alerts; it is not a lead,
        # except in test mode (#teste ... #fim), when it behaves as a new lead.
        command = "" if from_me else text_from_message(message or data).strip().lower()
        if command in {sdr.OWNER_TEST_START, sdr.OWNER_TEST_STOP}:
            if claim_message(external_id):
                await owner_test_command(jid, command)
            return {"received": True}
        if from_me or not sdr.owner_test_run(jid):
            return {"received": True}
    incoming = not from_me and event in {"MESSAGE", "MESSAGES_UPSERT", "BUTTONCLICK"}
    # Never name the recipient after the connected account's PushName.
    source_id, conversation_id = await ensure_conversation(jid, jid if from_me else str(push_name or jid))
    if not claim_message(external_id):
        return {"received": True}
    structured_event = payment_event(message or data, key=key, data=data)
    content = text_from_message(message or data)
    state_event = is_state_event(message or data) or event in {"REACTION", "MESSAGE_REACTION", "MESSAGE_EDIT", "MESSAGE_DELETE", "MESSAGE_REVOKE"}
    quoted_id, quoted_text, quoted_participant = quoted_context(message or data)
    if quoted_id and quoted_text:
        content = f'[Respondendo à mensagem de {quoted_participant or "contato"}: "{quoted_text}"]\n\n{content}'
    if incoming:
        # A password or token the lead pasted stops here: Chatwoot, Hermes and
        # the brain only ever see the mark (the batch worker sends the notice).
        content, _secret = sdr.redact_credentials(content)
    media = media_info(message or data)
    hermes_media = ([media] if media else []) + quoted_media_from_cache(message or data, quoted_id)
    if state_event:
        state_target_id = str((data.get("messageId") or data.get("targetId") or quoted_id or external_id) if isinstance(data, dict) else (quoted_id or external_id))
        record_event(external_id, event, jid, data if isinstance(data, dict) else {}, state_target_id)
        await sync_chatwoot_state(event, jid, state_target_id, data if isinstance(data, dict) else {})
    # Evolution can emit a first metadata-only copy and then a second copy
    # with base64. Do not consume the ID until the attachment is available.
    if has_media_envelope(message) and media is None and content == "[mídia recebida]":
        logger.info("Deferring media event until Evolution supplies base64 id=%s", external_id)
        return {"received": True}
    if incoming and media:
        try:
            remember_media(jid, external_id, media)
        except Exception:
            logger.exception("Failed to persist media history for jid=%s", jid)
    try:
        if from_me and not state_event:
            handover.record(conversation_id, external_id, "atendente_humano", content)
            await handover.pause(conversation_id)
        msg = await add_chatwoot_message(source_id, conversation_id, content, private=False, outgoing=not incoming, media=media, suppress_evolution=from_me, structured_event=structured_event)
        conn = db(); conn.execute("UPDATE messages SET chatwoot_id=? WHERE external_id=?", (msg.get("id"), external_id)); conn.commit(); conn.close()
    except Exception:
        conn = db(); conn.execute("DELETE FROM messages WHERE external_id=?", (external_id,)); conn.commit(); conn.close()
        raise
    if incoming and not state_event:
        queue_hermes_turn(jid, content, hermes_media or None, source_id, conversation_id, external_id, structured_event)
    return {"received": True}


@app.post("/chatwoot/webhook")
async def chatwoot_webhook(request: Request) -> dict[str, bool]:
    if BRIDGE_WEBHOOK_TOKEN and request.headers.get("authorization") != f"Bearer {BRIDGE_WEBHOOK_TOKEN}":
        return {"received": False}
    payload = await request.json()
    if payload.get("event") != "message_created" or payload.get("message_type") not in {"outgoing", 1, "1"}:
        return {"received": True}
    attributes = payload.get("content_attributes") or {}
    if isinstance(attributes, str):
        try:
            attributes = json.loads(attributes)
        except json.JSONDecodeError:
            attributes = {}
    if attributes.get("skip_evolution_bridge"):
        return {"received": True}
    delivery_id = str(payload.get("id") or "")
    if delivery_id:
        conn = db()
        try:
            conn.execute("BEGIN IMMEDIATE")
            inserted = conn.execute("INSERT OR IGNORE INTO chatwoot_deliveries VALUES (?)", (delivery_id,)).rowcount
            conn.commit()
        finally:
            conn.close()
        if not inserted:
            return {"received": True}
    content = payload.get("content") or ""
    attachments = payload.get("attachments") or []
    logger.info("Chatwoot outgoing id=%s content=%s attachments=%s", payload.get("id"), bool(content), len(attachments))
    if (not content and not attachments) or payload.get("private"):
        return {"received": True}
    conversation = payload.get("conversation") or {}
    conn = db(); row = conn.execute("SELECT jid FROM contacts WHERE conversation_id=?", (conversation.get("display_id") or conversation.get("id"),)).fetchone(); conn.close()
    if row:
        conversation_id = int(conversation.get("display_id") or conversation["id"])
        handover.record(conversation_id, "cw:" + delivery_id, "atendente_humano", content or "[anexo enviado pelo atendente]")
        await handover.pause(conversation_id)
        number = jid_number(row[0])
        native_sent = False
        native_location = attributes.get("whatsapp_location") or attributes.get("location")
        native_poll = attributes.get("whatsapp_poll") or attributes.get("poll")
        native_reaction = attributes.get("whatsapp_reaction") or attributes.get("reaction")
        if isinstance(native_location, dict):
            await send_evolution_location(number, native_location)
            native_sent = True
        elif isinstance(native_poll, dict):
            await send_evolution_poll(number, native_poll)
            native_sent = True
        elif isinstance(native_reaction, dict):
            await send_evolution_reaction(number, native_reaction)
            native_sent = True
        delivered_attachment = False
        for index, attachment in enumerate(attachments):
            media = await download_chatwoot_attachment(attachment)
            if media:
                await send_evolution_media(number, media, content if index == 0 else "")
                delivered_attachment = True
        if content and not delivered_attachment and not native_sent:
            await send_evolution(number, content)
    return {"received": True}
